"""
Energy-loss comparison between client-provided OCPP Excel and our CCTV sessions.

The Excel is authoritative for per-gun energy (`Units Consumed`). Our CCTV
sessions + MySQL meter are unreliable per-gun because:
  - Number plate (VRN) is often missing (Gemini Vision fails frequently)
  - Gun detection is noisy (slot_id → gun_number mapping is not trustworthy)
  - Meter only gives cumulative kWh across both guns

Strategy: weighted scoring across every signal available. Strong signals
(VRN, OCPP-start vs CCTV plug_time, duration) dominate when present; weak
signals (gun, model) keep pulling in the right direction when strong ones
are missing.

Gate: a candidate must clear the numeric score threshold AND have at least
one strong signal (exact VRN, a duration band, or a time band). Without
this, model + gun + same-day alone is enough to score 4 and produces
spurious matches between unrelated sessions of the same common model
(Tata TIAGO on Connector 1, etc.).

Excel "OCPP Start/End Time" — full timestamps when present. We compare
them to CCTV plug_time / in_time (both normalised to naive IST).
"""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any


# ── Config ──────────────────────────────────────────────────────────────────

WEIGHT_VRN_MATCH        = 3
# VRN-conflict penalty is 0 by design: user-typed VRNs on the OCPP side are
# noisy (typos, spacing variants, occasional wrong plates) and our CCTV-side
# VRN is LLM-extracted (also fallible). A mismatch is therefore NOT trustworthy
# negative evidence — only an exact normalized match is a positive signal.
# Strong matches still happen via time/duration/model; VRN can no longer veto.
WEIGHT_VRN_CONFLICT     = 0
WEIGHT_MODEL_MATCH      = 2
WEIGHT_GUN_MATCH        = 1
WEIGHT_DURATION_TIGHT   = 2   # within ±3 min
WEIGHT_DURATION_LOOSE   = 1   # within ±10 min
WEIGHT_DURATION_CONFLICT = -3 # > 10 min apart — likely different sessions
WEIGHT_TIME_TIGHT       = 3   # OCPP start within ±5 min of CCTV plug_time
WEIGHT_TIME_LOOSE       = 2   # within ±15 min
WEIGHT_TIME_CONFLICT    = -3  # > 30 min apart
WEIGHT_SAME_DATE        = 1
WEIGHT_DATE_CONFLICT    = -3

SCORE_THRESHOLD         = 3   # below → UNMATCHED

DURATION_TIGHT_SECONDS  = 3 * 60
DURATION_LOOSE_SECONDS  = 10 * 60
TIME_TIGHT_SECONDS      = 5 * 60
TIME_LOOSE_SECONDS      = 15 * 60
TIME_CONFLICT_SECONDS   = 30 * 60

# Station-specific slot→connector mapping. Flip this constant if wiring differs.
SLOT_TO_CONNECTOR: dict[str, int] = {
    "ROI_1": 1,
    "ROI_2": 2,
}


# ── Data classes ────────────────────────────────────────────────────────────

@dataclass
class ExcelRow:
    """One OCPP transaction from the client Excel (or a merged group thereof)."""
    row_index: int
    transaction_id: str | None
    session_id_ocpp: str | None
    start_dt: datetime | None      # OCPP Start Time (full timestamp, naive IST)
    end_dt: datetime | None        # OCPP End Time   (full timestamp, naive IST)
    duration_seconds: int | None   # from Session Duration(hh:mm:ss)
    connector_id: int | None
    vrn_raw: str | None
    vrn_norm: str | None
    make: str | None
    model: str | None
    units_kwh: float | None        # Units Consumed(kWh) — authoritative
    meter_start: float | None
    meter_end: float | None
    # Physical-bay grouping signals. id_tag + charge_point + connector_id is
    # the only reliable "same physical visit" key: VRN is user-typed and noisy,
    # model is LLM-extracted and hallucinates (Punch↔Tiago etc.), but id_tag
    # is the user's app/RFID identity and the bay is hardware.
    id_tag: str | None = None
    charge_point: str | None = None
    stop_reason: str | None = None   # Remote / EVDisconnected / ...
    closed_by: str | None = None     # balanceCutOff / mobile / CP / ...
    # Grouping bookkeeping. merged_count==1 means a raw OCPP row; >1 means
    # group_ocpp_transactions collapsed several adjacent rows into this one.
    merged_count: int = 1
    merged_transaction_ids: list[str] = field(default_factory=list)
    vrn_variants: list[str] = field(default_factory=list)   # distinct vrn_raw spellings in the group
    raw: dict[str, Any] = field(default_factory=dict)

    @property
    def date(self) -> datetime | None:
        """Day (midnight) of the OCPP start timestamp, for day-level pre-filter."""
        if self.start_dt is None:
            return None
        return self.start_dt.replace(hour=0, minute=0, second=0, microsecond=0)


@dataclass
class CctvSession:
    """Subset of our ChargingSession needed for matching."""
    session_id: int
    camera_id: str | None
    slot_id: str | None
    gun_number: str | None
    car_number: str | None
    car_model: str | None
    in_time: datetime | None
    plug_time: datetime | None
    plug_out_time: datetime | None
    out_time: datetime | None
    energy_kwh: float | None       # our meter-derived estimate

    @property
    def duration_seconds(self) -> int | None:
        """Prefer plug-based duration; fall back to car arrival/exit."""
        a, b = self.plug_time, self.plug_out_time
        if a is None or b is None:
            a, b = self.in_time, self.out_time
        if a is None or b is None:
            return None
        return int((b - a).total_seconds())

    @property
    def anchor_date(self) -> datetime | None:
        """The date this session belongs to (IST)."""
        t = self.plug_time or self.in_time
        if t is None:
            return None
        return _to_ist_date(t)


# ── Normalizers ─────────────────────────────────────────────────────────────

_VRN_STRIP = re.compile(r"[\s\-]+")

def normalize_vrn(v: Any) -> str | None:
    """KL 21 Z 1314 → KL21Z1314. Drops obvious junk like '1234'."""
    if v is None:
        return None
    s = str(v).strip().upper()
    s = _VRN_STRIP.sub("", s)
    if not s or s in {"-", "NA", "N/A", "NONE"}:
        return None
    # Plate must contain at least one letter to be considered valid for matching.
    if not any(c.isalpha() for c in s):
        return None
    return s


def normalize_model_text(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip().lower()
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s or None


def model_matches(cctv_model: str | None, excel_make: str | None, excel_model: str | None) -> bool:
    """Fuzzy: any token overlap between CCTV free-text and Excel Make+Model."""
    cctv = normalize_model_text(cctv_model)
    parts = normalize_model_text(f"{excel_make or ''} {excel_model or ''}")
    if not cctv or not parts:
        return False
    cctv_tokens = set(cctv.split())
    excel_tokens = set(parts.split())
    # Need at least one shared token ≥ 3 chars to avoid matching on "ev", "e".
    return any(t in cctv_tokens for t in excel_tokens if len(t) >= 3)


def parse_duration(s: Any) -> int | None:
    """hh:mm:ss → seconds. Also accepts timedelta (openpyxl sometimes gives that)."""
    if s is None:
        return None
    if isinstance(s, timedelta):
        return int(s.total_seconds())
    txt = str(s).strip()
    m = re.match(r"^(\d+):(\d+):(\d+)$", txt)
    if m:
        h, mnt, sec = map(int, m.groups())
        return h * 3600 + mnt * 60 + sec
    return None


def parse_ocpp_datetime(v: Any) -> datetime | None:
    """Excel 'OCPP Start/End Time' — accepts datetime objects (openpyxl) or
    strings like 'dd/mm/yyyy HH:MM:SS' / 'dd/mm/yyyy'. Returns a naive datetime
    in IST (the client export is already IST)."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.replace(tzinfo=None) if v.tzinfo else v
    txt = str(v).strip()
    fmts = (
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
        "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y",
    )
    for fmt in fmts:
        try:
            return datetime.strptime(txt, fmt)
        except ValueError:
            continue
    return None


def _to_ist_naive(dt: datetime | None) -> datetime | None:
    """Normalise a (possibly tz-aware UTC) datetime to a naive IST datetime."""
    if dt is None:
        return None
    from zoneinfo import ZoneInfo
    if dt.tzinfo is not None:
        dt = dt.astimezone(ZoneInfo("Asia/Kolkata")).replace(tzinfo=None)
    return dt


def _to_ist_date(dt: datetime) -> datetime:
    """Date of `dt` expressed in IST, as a naive midnight datetime for comparison."""
    naive = _to_ist_naive(dt)
    if naive is None:
        return None  # type: ignore[return-value]
    return naive.replace(hour=0, minute=0, second=0, microsecond=0)


# ── Parser ──────────────────────────────────────────────────────────────────

# Header → ExcelRow field. Tolerates capitalization / whitespace drift.
_HEADER_MAP = {
    "transaction id":        "transaction_id",
    "session id":            "session_id_ocpp",
    "ocpp start time":       "_start",
    "ocpp end time":         "_end",
    "session duration":      "_duration",          # matches "Session Duration(hh:mm:ss)"
    "connector id":          "_connector",
    "vrn":                   "_vrn",
    "make":                  "make",
    "model":                 "model",
    "units consumed":        "_units",             # matches "Units Consumed(kWh)"
    "meter start":           "_meter_start",
    "meter end":             "_meter_end",
    "id tag":                "_id_tag",            # user identity — grouping key
    "charge point":          "_charge_point",      # physical station — grouping key
    "stop reason":           "_stop_reason",       # Remote / EVDisconnected
    "closed by":             "_closed_by",         # balanceCutOff / mobile / CP
}


def _header_key(h: Any) -> str:
    """Strip parenthetical suffix so 'Units Consumed(kWh)' → 'units consumed'."""
    if h is None:
        return ""
    s = str(h).strip().lower()
    s = re.sub(r"\s*\(.*?\)\s*", "", s)   # drop (kWh), (hh:mm:ss)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def parse_excel(file_bytes: bytes) -> list[ExcelRow]:
    """Parse the client OCPP export xlsx into ExcelRow records."""
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    headers = next(rows_iter, None)
    if not headers:
        return []
    col_idx: dict[str, int] = {}
    for i, h in enumerate(headers):
        key = _header_key(h)
        if key in _HEADER_MAP:
            col_idx[_HEADER_MAP[key]] = i

    parsed: list[ExcelRow] = []
    for row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or c == "" for c in row):
            continue

        def get(field_name: str):
            i = col_idx.get(field_name)
            if i is None or i >= len(row):
                return None
            return row[i]

        vrn_raw = get("_vrn")
        duration_s = parse_duration(get("_duration"))
        start_dt = parse_ocpp_datetime(get("_start"))
        end_dt = parse_ocpp_datetime(get("_end"))
        connector_raw = get("_connector")
        try:
            connector_id = int(connector_raw) if connector_raw not in (None, "") else None
        except (TypeError, ValueError):
            connector_id = None

        def _to_float(v):
            if v in (None, "", "-"):
                return None
            try:
                return float(v)
            except (TypeError, ValueError):
                return None

        def _str_or_none(v):
            if v is None:
                return None
            s = str(v).strip()
            return s or None

        parsed.append(ExcelRow(
            row_index=row_num,
            transaction_id=str(get("transaction_id")) if get("transaction_id") is not None else None,
            session_id_ocpp=str(get("session_id_ocpp")) if get("session_id_ocpp") is not None else None,
            start_dt=start_dt,
            end_dt=end_dt,
            duration_seconds=duration_s,
            connector_id=connector_id,
            vrn_raw=str(vrn_raw).strip() if vrn_raw else None,
            vrn_norm=normalize_vrn(vrn_raw),
            make=str(get("make")).strip() if get("make") else None,
            model=str(get("model")).strip() if get("model") else None,
            units_kwh=_to_float(get("_units")),
            meter_start=_to_float(get("_meter_start")),
            meter_end=_to_float(get("_meter_end")),
            id_tag=_str_or_none(get("_id_tag")),
            charge_point=_str_or_none(get("_charge_point")),
            stop_reason=_str_or_none(get("_stop_reason")),
            closed_by=_str_or_none(get("_closed_by")),
        ))
    return parsed


# ── Physical-session grouping ───────────────────────────────────────────────
#
# Why this exists:
# A single physical parking session frequently produces multiple OCPP rows.
# A user runs out of balance mid-charge (Stop Reason="Remote",
# Closed By="balanceCutOff"), tops up via the app, and resumes within minutes.
# Each restart is a fresh OCPP transaction but the car never moved. The CCTV
# side sees one continuous ChargingSession, so per-row matching strands the
# extra OCPP rows as "No match" and the matched one carries only a fraction
# of the real kWh — making the loss number meaningless.
#
# The reliable grouping signal is the physical bay + the user identity:
#   - VRN: user-typed, missing/typo'd often → cannot trust
#   - Make/Model: extracted by LLM on the CCTV side, hallucinates between
#     similar models (Tata Punch ↔ Tata Tiago) → cannot trust
#   - id_tag + charge_point + connector_id: stable hardware/account identity
#     → trust
#
# Plus two adjacency checks that must hold for a merge:
#   - Time gap between rows < GROUP_TIME_GAP_MAX_SECONDS (came back quickly)
#   - Meter continuity: B.meter_start ≈ A.meter_end (same physical plug)

# Time cap is the sanity ceiling; meter continuity is the real proof. A
# < 200 Wh drift between A.meter_end and B.meter_start means the connector
# physically never released the car. 4h covers extended balance-top-up
# breaks observed in the 2026-05-10 Volvo data (53 min + 1h 48m gaps with
# 47 / 58 Wh meter drift across them).
GROUP_TIME_GAP_MAX_SECONDS = 4 * 60 * 60   # 4 hours
GROUP_METER_DRIFT_MAX_WH   = 200           # 0.2 kWh
# Meter readings in this dataset are in Wh (e.g. 54336429 → 54348238 = 11.81 kWh).


def _has_group_keys(r: ExcelRow) -> bool:
    """A row can participate in grouping only when all physical keys are present."""
    return (
        r.id_tag is not None
        and r.charge_point is not None
        and r.connector_id is not None
        and r.start_dt is not None
    )


def _can_merge(prev: ExcelRow, curr: ExcelRow) -> bool:
    """True when prev and curr are adjacent OCPP rows of the same physical session."""
    if (prev.charge_point, prev.connector_id, prev.id_tag) != \
       (curr.charge_point, curr.connector_id, curr.id_tag):
        return False
    if prev.end_dt is None or curr.start_dt is None:
        return False
    gap_s = (curr.start_dt - prev.end_dt).total_seconds()
    # Allow tiny negative drift (clock skew between OCPP server and charger)
    # but reject real overlap — overlapping transactions on one connector are
    # a data-quality problem, not a same-session signal.
    if gap_s < -60 or gap_s >= GROUP_TIME_GAP_MAX_SECONDS:
        return False
    if prev.meter_end is None or curr.meter_start is None:
        return False
    meter_diff = curr.meter_start - prev.meter_end
    # Meter is monotonic on a connector — non-negative diff only, within tolerance.
    return 0 <= meter_diff < GROUP_METER_DRIFT_MAX_WH


def _merge_group(group: list[ExcelRow]) -> ExcelRow:
    """Collapse a list of adjacent OCPP rows into one merged ExcelRow."""
    if len(group) == 1:
        r = group[0]
        # Even singletons carry their transaction_id in the merged list so
        # the dashboard's drill-down has a uniform shape.
        if not r.merged_transaction_ids and r.transaction_id:
            r.merged_transaction_ids = [r.transaction_id]
        if r.vrn_raw and not r.vrn_variants:
            r.vrn_variants = [r.vrn_raw]
        return r

    first, last = group[0], group[-1]

    def _mode(values: list[Any]) -> Any:
        clean = [v for v in values if v not in (None, "")]
        if not clean:
            return None
        return Counter(clean).most_common(1)[0][0]

    # Sums (None when every row is None for that field).
    units_vals    = [r.units_kwh        for r in group if r.units_kwh        is not None]
    duration_vals = [r.duration_seconds for r in group if r.duration_seconds is not None]
    total_units    = sum(units_vals)    if units_vals    else None
    total_duration = sum(duration_vals) if duration_vals else None

    # Mode-by-frequency on the normalized VRN (typo-tolerant). Pick a raw
    # spelling that maps back to the mode; on tie, prefer the longest one
    # so the dashboard surfaces the most plate detail (e.g. "KL 64 L 5395"
    # over "KL64L5395").
    vrn_norm_mode = _mode([r.vrn_norm for r in group])
    raw_candidates = [r.vrn_raw for r in group if r.vrn_norm == vrn_norm_mode and r.vrn_raw]
    vrn_raw_mode = max(raw_candidates, key=len) if raw_candidates else None

    # Distinct raw spellings (input order preserved) — exposes user-typed
    # variation for the dashboard / audit trail.
    seen: dict[str, None] = {}
    for r in group:
        if r.vrn_raw and r.vrn_raw not in seen:
            seen[r.vrn_raw] = None
    vrn_variants = list(seen.keys())

    tx_ids = [r.transaction_id for r in group if r.transaction_id]

    return ExcelRow(
        row_index=first.row_index,
        transaction_id=first.transaction_id,     # canonical = first
        session_id_ocpp=first.session_id_ocpp,
        start_dt=min((r.start_dt for r in group if r.start_dt is not None), default=None),
        end_dt  =max((r.end_dt   for r in group if r.end_dt   is not None), default=None),
        duration_seconds=total_duration,
        connector_id=first.connector_id,
        vrn_raw=vrn_raw_mode,
        vrn_norm=vrn_norm_mode,
        make=_mode([r.make  for r in group]),
        model=_mode([r.model for r in group]),
        units_kwh=total_units,
        meter_start=first.meter_start,
        meter_end=last.meter_end,
        id_tag=first.id_tag,
        charge_point=first.charge_point,
        # Last row's stop reason is the real exit signal. Intermediate
        # balanceCutOffs are interruptions, not the final termination.
        stop_reason=last.stop_reason,
        closed_by=last.closed_by,
        merged_count=len(group),
        merged_transaction_ids=tx_ids,
        vrn_variants=vrn_variants,
    )


def group_ocpp_transactions(rows: list[ExcelRow]) -> list[ExcelRow]:
    """
    Collapse adjacent OCPP transactions of the same physical session into one.

    See module-level comment block above for the rule and rationale. Rows
    that lack any grouping key (id_tag / charge_point / connector_id /
    start_dt) pass through unchanged — they can't participate in a merge
    safely without those signals.

    Output preserves the original row_index ordering of the file so the
    dashboard sees a stable layout.
    """
    if not rows:
        return rows

    sortable = [r for r in rows if _has_group_keys(r)]
    passthrough = [r for r in rows if not _has_group_keys(r)]

    sortable.sort(key=lambda r: (
        r.charge_point or "",
        r.connector_id if r.connector_id is not None else -1,
        r.id_tag or "",
        r.start_dt or datetime.min,
    ))

    groups: list[list[ExcelRow]] = []
    for r in sortable:
        if groups and _can_merge(groups[-1][-1], r):
            groups[-1].append(r)
        else:
            groups.append([r])

    merged = [_merge_group(grp) for grp in groups]

    out = merged + passthrough
    out.sort(key=lambda r: r.row_index)
    return out


# ── Matcher ─────────────────────────────────────────────────────────────────

@dataclass
class MatchResult:
    excel_row: ExcelRow
    cctv_session: CctvSession | None
    score: int
    reasons: list[str]
    confidence: str  # HIGH | MEDIUM | LOW | UNMATCHED

    @property
    def client_kwh(self) -> float | None:
        return self.excel_row.units_kwh

    @property
    def meter_kwh(self) -> float | None:
        return self.cctv_session.energy_kwh if self.cctv_session else None

    @property
    def loss_kwh(self) -> float | None:
        c, m = self.client_kwh, self.meter_kwh
        if c is None or m is None:
            return None
        return round(c - m, 3)

    @property
    def loss_pct(self) -> float | None:
        c, l = self.client_kwh, self.loss_kwh
        if c is None or l is None or c == 0:
            return None
        return round((l / c) * 100, 1)


def _score_pair(excel: ExcelRow, cctv: CctvSession) -> tuple[int, list[str]]:
    score = 0
    reasons: list[str] = []

    # Date — pre-filter but also scored for clarity in the reasons trail.
    excel_date = excel.date
    cctv_date = cctv.anchor_date
    if excel_date and cctv_date:
        if excel_date == cctv_date:
            score += WEIGHT_SAME_DATE
            reasons.append("date")
        else:
            score += WEIGHT_DATE_CONFLICT
            reasons.append("date-conflict")

    # VRN — strongest signal when both present and valid.
    if excel.vrn_norm and cctv.car_number:
        cctv_vrn = normalize_vrn(cctv.car_number)
        if cctv_vrn:
            if cctv_vrn == excel.vrn_norm:
                score += WEIGHT_VRN_MATCH
                reasons.append("VRN")
            else:
                score += WEIGHT_VRN_CONFLICT
                reasons.append("VRN-conflict")

    # Model — fuzzy token overlap.
    if model_matches(cctv.car_model, excel.make, excel.model):
        score += WEIGHT_MODEL_MATCH
        reasons.append("model")

    # Gun ↔ connector.
    if excel.connector_id is not None and cctv.slot_id is not None:
        expected_connector = SLOT_TO_CONNECTOR.get(cctv.slot_id)
        if expected_connector is not None and expected_connector == excel.connector_id:
            score += WEIGHT_GUN_MATCH
            reasons.append("gun")

    # Duration proximity (or penalty when both are known and far apart).
    # Use the OCPP wall-clock span (end_dt - start_dt) rather than the
    # reported duration_seconds. For ungrouped rows these are equal. For
    # grouped rows, duration_seconds is the SUM of charging time across N
    # transactions (e.g. 1h 55m for the Volvo's 3 sub-sessions) while the
    # wall-clock span is start-of-first to end-of-last (4h 36m) — which is
    # what lines up with CCTV's plug_time → plug_out_time window.
    cctv_dur = cctv.duration_seconds
    excel_wall_dur = None
    if excel.start_dt is not None and excel.end_dt is not None:
        excel_wall_dur = int((excel.end_dt - excel.start_dt).total_seconds())
    if excel_wall_dur is not None and cctv_dur is not None:
        diff = abs(excel_wall_dur - cctv_dur)
        if diff <= DURATION_TIGHT_SECONDS:
            score += WEIGHT_DURATION_TIGHT
            reasons.append(f"duration±3m({diff}s)")
        elif diff <= DURATION_LOOSE_SECONDS:
            score += WEIGHT_DURATION_LOOSE
            reasons.append(f"duration±10m({diff}s)")
        else:
            score += WEIGHT_DURATION_CONFLICT
            reasons.append(f"duration-conflict({diff}s)")

    # Wall-clock proximity: OCPP start vs CCTV plug_time (fallback in_time).
    # Strongest single non-VRN signal — same gun at the same minute is decisive.
    excel_start = _to_ist_naive(excel.start_dt)
    cctv_anchor = _to_ist_naive(cctv.plug_time) or _to_ist_naive(cctv.in_time)
    if excel_start is not None and cctv_anchor is not None:
        time_diff = abs(int((excel_start - cctv_anchor).total_seconds()))
        if time_diff <= TIME_TIGHT_SECONDS:
            score += WEIGHT_TIME_TIGHT
            reasons.append(f"time±5m({time_diff}s)")
        elif time_diff <= TIME_LOOSE_SECONDS:
            score += WEIGHT_TIME_LOOSE
            reasons.append(f"time±15m({time_diff}s)")
        elif time_diff > TIME_CONFLICT_SECONDS:
            score += WEIGHT_TIME_CONFLICT
            reasons.append(f"time-conflict({time_diff}s)")

    return score, reasons


def _has_strong_signal(reasons: list[str]) -> bool:
    """A candidate must clear at least one of: exact VRN, a duration band, or
    a time band. Without any of these, model+gun+date alone is too weak —
    that combination produced the spurious LOW matches in the client sample."""
    if "VRN" in reasons:
        return True
    for r in reasons:
        if r.startswith("duration±") or r.startswith("time±"):
            return True
    return False


def _confidence(score: int, reasons: list[str]) -> str:
    if score < SCORE_THRESHOLD:
        return "UNMATCHED"
    has_vrn = "VRN" in reasons
    has_time_tight = any(r.startswith("time±5m") for r in reasons)
    has_dur_tight = any(r.startswith("duration±3m") for r in reasons)
    if has_vrn and (has_time_tight or has_dur_tight):
        return "HIGH"
    if has_time_tight and has_dur_tight:
        return "HIGH"
    if score >= 6:
        return "HIGH"
    if score >= 4:
        return "MEDIUM"
    return "LOW"


def match_excel_to_cctv(excel_rows: list[ExcelRow], cctv_sessions: list[CctvSession]) -> list[MatchResult]:
    """Greedy highest-score-first assignment.

    Each CCTV session is consumed at most once. If no candidate clears the
    SCORE_THRESHOLD for an Excel row, it is reported as UNMATCHED (typical for
    blind-spot parking or a camera outage during that session).
    """
    # Build candidate scores (excel_index, cctv_index, score, reasons).
    candidates: list[tuple[int, int, int, list[str]]] = []
    for ei, ex in enumerate(excel_rows):
        for ci, cc in enumerate(cctv_sessions):
            score, reasons = _score_pair(ex, cc)
            # Two-stage filter: numeric threshold + at least one strong signal.
            # Without the strong-signal gate, common-model + same-gun + same-day
            # bundles score 4 and produce nonsense pairings (e.g. an 11-second
            # OCPP cancel matched to an unrelated half-hour TIAGO session).
            if score >= SCORE_THRESHOLD and _has_strong_signal(reasons):
                candidates.append((ei, ci, score, reasons))

    candidates.sort(key=lambda t: (-t[2], t[0]))

    used_excel: set[int] = set()
    used_cctv: set[int] = set()
    results_by_ei: dict[int, MatchResult] = {}

    for ei, ci, score, reasons in candidates:
        if ei in used_excel or ci in used_cctv:
            continue
        used_excel.add(ei)
        used_cctv.add(ci)
        results_by_ei[ei] = MatchResult(
            excel_row=excel_rows[ei],
            cctv_session=cctv_sessions[ci],
            score=score,
            reasons=reasons,
            confidence=_confidence(score, reasons),
        )

    # Emit one result per Excel row, in original order.
    results: list[MatchResult] = []
    for ei, ex in enumerate(excel_rows):
        if ei in results_by_ei:
            results.append(results_by_ei[ei])
        else:
            results.append(MatchResult(
                excel_row=ex,
                cctv_session=None,
                score=0,
                reasons=[],
                confidence="UNMATCHED",
            ))
    return results


def result_to_dict(r: MatchResult) -> dict[str, Any]:
    """Serializable shape for the dashboard."""
    ex = r.excel_row
    cc = r.cctv_session
    return {
        "excel_row": ex.row_index,
        "transaction_id": ex.transaction_id,
        "ocpp_date": ex.date.strftime("%Y-%m-%d") if ex.date else None,
        "ocpp_start_time": ex.start_dt.strftime("%Y-%m-%d %H:%M:%S") if ex.start_dt else None,
        "ocpp_end_time":   ex.end_dt.strftime("%Y-%m-%d %H:%M:%S")   if ex.end_dt   else None,
        "duration_seconds": ex.duration_seconds,
        "connector_id": ex.connector_id,
        "vrn": ex.vrn_raw,
        "make": ex.make,
        "model": ex.model,
        "client_kwh": ex.units_kwh,
        "meter_kwh": r.meter_kwh,
        "loss_kwh": r.loss_kwh,
        "loss_pct": r.loss_pct,
        "matched_session_id": cc.session_id if cc else None,
        "matched_slot_id": cc.slot_id if cc else None,
        "matched_car_number": cc.car_number if cc else None,
        "matched_car_model": cc.car_model if cc else None,
        "match_score": r.score,
        "match_reasons": r.reasons,
        "confidence": r.confidence,
        # Physical-session grouping metadata. merged_count > 1 means
        # group_ocpp_transactions collapsed N raw OCPP rows into this one;
        # the dashboard can show a "merged from N" badge and drill down via
        # transaction_ids. vrn_variants lists every distinct user-typed
        # spelling we saw in the group (typo evidence).
        "merged_count": ex.merged_count,
        "transaction_ids": ex.merged_transaction_ids,
        "vrn_variants": ex.vrn_variants,
        "stop_reason": ex.stop_reason,
        "closed_by": ex.closed_by,
        "id_tag": ex.id_tag,
        "charge_point": ex.charge_point,
    }


# ── Charging-deviation detection ─────────────────────────────────────────────
#
# "Flag sessions whose energy is abnormal for that car model, track the
# specific vehicle, and see if the same car deviates elsewhere."
#
# We baseline off the CLIENT (OCPP Excel) side ONLY. units_kwh is the
# authoritative meter of record and make/model/vrn/duration/charge_point all
# come from the billing system — so a deviation flag never depends on a CCTV
# match landing. (Our CCTV meter estimate is an equal-share allocation of a
# shared physical meter and is far too noisy to baseline against.)
#
# Baseline = per-car-model robust statistics. We flag on TWO axes because raw
# kWh conflates with how long the car charged:
#   - energy : units_kwh            vs the model's typical session energy
#   - rate   : units_kwh / hours    vs the model's typical charging rate
# A session is flagged if EITHER axis exceeds the cutoff. Rate is the sharper
# "consumed so much energy" signal — a car pulling far more kWh per hour than
# its model normally does is the real anomaly, independent of session length.
#
# Robust stats (median + MAD, Iglewicz-Hoaglin modified z) rather than
# mean/std: the per-model sample is small and itself contains the outliers we
# are hunting, so a single 60 kWh session would inflate a std-based band
# enough to hide itself. Models with too few samples for their own baseline
# fall back to the GLOBAL rate distribution so those cars are never silently
# skipped (we record which baseline was used in the reasons trail).
#
# Vehicle tracking: deviations are rolled up by normalized VRN. The rollup
# also carries the distinct charge_points a plate appeared on — so when a
# single upload spans multiple stations, a recurring deviant is visible
# ACROSS stations on the billing side today. (CCTV-side cross-station
# tracking still needs reliable ANPR + multi-station wiring — out of scope
# here.)

DEVIATION_Z_THRESHOLD  = 3.5        # Iglewicz-Hoaglin standard cutoff
DEVIATION_MIN_SAMPLES  = 4          # need a real baseline before flagging a model
DEVIATION_MIN_HOURS    = 1.0 / 60   # 1 min floor — sub-minute rows give a garbage rate


def _median(xs: list[float]) -> float | None:
    s = sorted(xs)
    n = len(s)
    if n == 0:
        return None
    mid = n // 2
    return s[mid] if n % 2 else (s[mid - 1] + s[mid]) / 2.0


def _modified_z(x: float, med: float | None, mad: float | None, meanad: float | None) -> float | None:
    """Iglewicz-Hoaglin modified z-score. Falls back to a mean-abs-deviation
    scale when MAD is 0 (e.g. a run of identical values with one outlier),
    so the outlier still flags instead of dividing by zero."""
    if med is None:
        return None
    if mad and mad > 0:
        return 0.6745 * (x - med) / mad
    if meanad and meanad > 0:
        return (x - med) / (1.253314 * meanad)
    return None


def _baseline(values: list[float]) -> dict[str, Any] | None:
    """Median / MAD / mean-abs-deviation for one metric of one group."""
    if not values:
        return None
    med = _median(values)
    devs = [abs(v - med) for v in values]
    mad = _median(devs)
    meanad = sum(devs) / len(devs) if devs else 0.0
    return {
        "samples": len(values),
        "median": round(med, 3),
        "mad": round(mad, 3) if mad is not None else None,
        "_med": med, "_mad": mad, "_meanad": meanad,   # unrounded, for scoring
    }


def _model_key(make: str | None, model: str | None) -> str:
    """Group key for a car model, tolerant of case / spacing. Matches the
    make+model label used by the energy-analysis endpoint."""
    norm = normalize_model_text(f"{make or ''} {model or ''}")
    return norm or "unknown"


def _model_label(make: str | None, model: str | None) -> str:
    return " ".join(filter(None, [make, model])).strip() or "Unknown"


def _axis_fires(z: float | None, threshold: float, direction: str) -> bool:
    """Whether an axis's modified-z clears the threshold in the wanted
    direction. 'high' = over-consumption only (z>0), 'low' = under only,
    'both' = either side."""
    if z is None or abs(z) < threshold:
        return False
    if direction == "high":
        return z > 0
    if direction == "low":
        return z < 0
    return True


def flag_energy_deviations(
    results: list[dict[str, Any]],
    *,
    z_threshold: float = DEVIATION_Z_THRESHOLD,
    min_samples: int = DEVIATION_MIN_SAMPLES,
    direction: str = "high",
) -> dict[str, Any]:
    """
    Flag charging sessions whose energy deviates from normal for their car
    model, and roll deviations up per vehicle.

    Input: the `results` list produced by result_to_dict (one dict per OCPP
    row, matched or not — we only read the authoritative Excel-side fields:
    client_kwh, make, model, vrn, duration_seconds, connector_id,
    charge_point, transaction_id).

    direction — which deviations to flag:
      "high" (default) : only OVER-consumption (abnormally high energy/rate),
                         matching the "a car that consumed so much energy"
                         goal. Drops slow-charge noise; also stops a tight
                         small-sample model group from flagging a session
                         purely for charging slightly slower than usual.
      "low"            : only under-consumption / slow charges.
      "both"           : either side.
    Both z-scores are always reported for transparency; only axes that fire
    in the wanted direction drive `reasons`, `severity`, and the count.

    Returns { baselines, deviations, vehicles, summary }.
    """
    direction = (direction or "high").strip().lower()
    if direction not in ("high", "low", "both"):
        direction = "high"
    # ── Considered rows: authoritative energy present and positive ────────────
    considered: list[dict[str, Any]] = []
    for r in results:
        kwh = r.get("client_kwh")
        if kwh is None or kwh <= 0:
            continue
        dur_s = r.get("duration_seconds")
        hours = (dur_s / 3600.0) if isinstance(dur_s, (int, float)) and dur_s else None
        rate = (kwh / hours) if (hours is not None and hours >= DEVIATION_MIN_HOURS) else None
        considered.append({
            "row": r,
            "mkey": _model_key(r.get("make"), r.get("model")),
            "label": _model_label(r.get("make"), r.get("model")),
            "kwh": float(kwh),
            "hours": round(hours, 3) if hours is not None else None,
            "rate": round(rate, 3) if rate is not None else None,
            "_rate": rate,
        })

    # ── Per-model baselines (energy + rate) and a global rate fallback ────────
    by_model: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for c in considered:
        by_model[c["mkey"]].append(c)

    model_baselines: dict[str, dict[str, Any]] = {}
    for mkey, group in by_model.items():
        kwh_base = _baseline([c["kwh"] for c in group])
        rate_base = _baseline([c["_rate"] for c in group if c["_rate"] is not None])
        model_baselines[mkey] = {
            "model": group[0]["label"],
            "sessions": len(group),
            "energy": kwh_base,
            "rate": rate_base,
        }

    global_rate = _baseline([c["_rate"] for c in considered if c["_rate"] is not None])

    # ── Score each session on both axes ───────────────────────────────────────
    deviations: list[dict[str, Any]] = []
    for c in considered:
        mb = model_baselines[c["mkey"]]
        reasons: list[str] = []
        kwh_z = rate_z = None
        firing: list[float] = []   # signed z of axes that fired in-direction

        # Energy axis — only within a model that has a real baseline.
        eb = mb["energy"]
        if eb and eb["samples"] >= min_samples:
            kwh_z = _modified_z(c["kwh"], eb["_med"], eb["_mad"], eb["_meanad"])
            if _axis_fires(kwh_z, z_threshold, direction):
                reasons.append(f"energy {'high' if kwh_z > 0 else 'low'} vs {mb['model']} "
                               f"(z={kwh_z:.1f}, median={eb['median']} kWh)")
                firing.append(kwh_z)

        # Rate axis — prefer the model baseline; fall back to global.
        if c["_rate"] is not None:
            rb = mb["rate"]
            if rb and rb["samples"] >= min_samples:
                rate_z = _modified_z(c["_rate"], rb["_med"], rb["_mad"], rb["_meanad"])
                base_lbl, base_med = mb["model"], rb["median"]
            elif global_rate and global_rate["samples"] >= min_samples:
                rate_z = _modified_z(c["_rate"], global_rate["_med"], global_rate["_mad"], global_rate["_meanad"])
                base_lbl, base_med = "all models", global_rate["median"]
            else:
                base_lbl = base_med = None
            if _axis_fires(rate_z, z_threshold, direction):
                reasons.append(f"rate {'high' if rate_z > 0 else 'low'} vs {base_lbl} "
                               f"(z={rate_z:.1f}, median={base_med} kWh/h)")
                firing.append(rate_z)

        if not reasons:
            continue

        r = c["row"]
        # Severity and direction come from the FIRING axes only, so an
        # out-of-direction axis (e.g. a slow-charge rate when direction="high")
        # never inflates severity or mislabels the row.
        dominant = max(firing, key=abs)
        severity = abs(dominant)
        deviations.append({
            "transaction_id": r.get("transaction_id"),
            "excel_row": r.get("excel_row"),
            "ocpp_start_time": r.get("ocpp_start_time"),
            "vrn": r.get("vrn"),
            "vrn_norm": normalize_vrn(r.get("vrn")),
            "make": r.get("make"),
            "model": r.get("model"),
            "model_label": c["label"],
            "connector_id": r.get("connector_id"),
            "charge_point": r.get("charge_point"),
            "client_kwh": round(c["kwh"], 3),
            "duration_hours": c["hours"],
            "charge_rate_kwh_h": c["rate"],
            "energy_z": round(kwh_z, 2) if kwh_z is not None else None,
            "rate_z": round(rate_z, 2) if rate_z is not None else None,
            "direction": "high" if dominant > 0 else "low",
            "severity": round(severity, 2),
            "reasons": reasons,
            # Carry the CCTV match through so the operator can jump to footage.
            "matched_session_id": r.get("matched_session_id"),
            "matched_car_number": r.get("matched_car_number"),
            "confidence": r.get("confidence"),
        })

    deviations.sort(key=lambda d: -d["severity"])
    flagged_txn = {d["transaction_id"] for d in deviations if d["transaction_id"]}

    # ── Per-vehicle rollup — track the specific car across the dataset ────────
    # Keyed on normalized VRN (the only cross-session car handle we have).
    # `stations` exposes cross-station recurrence on the billing side.
    veh: dict[str, dict[str, Any]] = {}
    for c in considered:
        r = c["row"]
        vnorm = normalize_vrn(r.get("vrn"))
        if not vnorm:
            continue
        v = veh.setdefault(vnorm, {
            "vrn_norm": vnorm,
            "vrn": r.get("vrn"),
            "sessions": 0,
            "deviations": 0,
            "stations": set(),
            "connectors": set(),
            "models": set(),
            "deviation_txns": [],
        })
        v["sessions"] += 1
        if r.get("charge_point"):
            v["stations"].add(r["charge_point"])
        if r.get("connector_id") is not None:
            v["connectors"].add(r["connector_id"])
        if c["label"] != "Unknown":
            v["models"].add(c["label"])
        if r.get("transaction_id") in flagged_txn:
            v["deviations"] += 1
            v["deviation_txns"].append(r.get("transaction_id"))

    vehicles = []
    for v in veh.values():
        if v["deviations"] == 0:
            continue   # only surface cars that actually deviated
        vehicles.append({
            "vrn": v["vrn"],
            "vrn_norm": v["vrn_norm"],
            "sessions": v["sessions"],
            "deviations": v["deviations"],
            "recurring": v["deviations"] >= 2,           # deviated more than once
            "multi_station": len(v["stations"]) > 1,     # same car, different stations
            "stations": sorted(v["stations"]),
            "connectors": sorted(v["connectors"]),
            "models": sorted(v["models"]),
            "deviation_txns": v["deviation_txns"],
        })
    vehicles.sort(key=lambda x: (-x["deviations"], -x["sessions"]))

    # ── Serializable baseline table (drop the unrounded scoring internals) ────
    baseline_table = []
    for mkey, mb in model_baselines.items():
        def _clean(b):
            return {k: val for k, val in b.items() if not k.startswith("_")} if b else None
        baseline_table.append({
            "model": mb["model"],
            "sessions": mb["sessions"],
            "energy": _clean(mb["energy"]),
            "rate": _clean(mb["rate"]),
            "baselined": bool(mb["energy"] and mb["energy"]["samples"] >= min_samples),
        })
    baseline_table.sort(key=lambda x: -x["sessions"])

    return {
        "baselines": {
            "by_model": baseline_table,
            "global_rate": {k: val for k, val in global_rate.items() if not k.startswith("_")} if global_rate else None,
            "z_threshold": z_threshold,
            "min_samples": min_samples,
            "direction": direction,
        },
        "deviations": deviations,
        "vehicles": vehicles,
        "summary": {
            "considered": len(considered),
            "flagged": len(deviations),
            "vehicles_flagged": len(vehicles),
            "recurring_vehicles": sum(1 for v in vehicles if v["recurring"]),
            "multi_station_vehicles": sum(1 for v in vehicles if v["multi_station"]),
        },
    }
